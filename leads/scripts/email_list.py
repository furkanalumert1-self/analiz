"""Step 8b: one-row-per-email contact list (reports/callypso_email_list.xlsx). Run after build.py, before render_page.py."""
import json, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

H = os.path.dirname(__file__)
OUT = os.path.join(H, "..", "..", "reports")
C = os.path.join(H, "..", "cache")
d = json.load(open(os.path.join(OUT, "leads-data.json")))
CE = json.load(open(os.path.join(C, "company_emails.json")))
FREE = re.compile(r"@(gmail|googlemail|hotmail|outlook|live|yahoo|icloud|gmx|web)\.", re.I)
SKIP = re.compile(r"^(nfo|bewerbung|jobs?|karriere|careers?|privacy|datenschutz)@")

COLS = ["Company Name", "Website", "Country", "City", "Industry", "Priority", "Engage Fit", "Contact Type", "Contact Name",
        "Job Title", "Email", "Email Domain Type", "Email Verification / Source", "Source URL", "LinkedIn URL"]
rows = []
for r in d["rows"]:
    base = {k: r[k] for k in ("Company Name", "Website", "Country", "City", "Industry", "Priority", "Engage Fit")}
    for p in ("Primary", "Secondary"):
        if r.get(f"{p} Email"):
            src = r[f"{p} Email Source"]
            rows.append(dict(base, **{"Contact Type": f"Decision maker ({p.lower()})", "Contact Name": r[f"{p} Decision Maker"],
                                      "Job Title": r[f"{p} Job Title"], "Email": r[f"{p} Email"],
                                      "Email Verification / Source": src.split(" (http")[0],
                                      "Source URL": (re.search(r"\((https?://[^)]+)\)", src) or [None, ""])[1],
                                      "LinkedIn URL": r[f"{p} LinkedIn URL"]}))
    person = {x["Email"] for x in rows}
    dom = re.sub(r"^https?://(www\.)?", "", r["Website"]).split("/")[0]
    for key in CE:
        if dom.endswith(key):
            for e, u in CE[key]["emails"]:
                if e in person or SKIP.search(e):
                    continue
                rows.append(dict(base, **{"Contact Type": "Company (general)", "Contact Name": "", "Job Title": "", "Email": e,
                                          "Email Verification / Source": "Published on company website", "Source URL": u,
                                          "LinkedIn URL": r["LinkedIn Company URL"]}))
            break
for x in rows:
    x["Email Domain Type"] = "Free mail (not company domain)" if FREE.search(x["Email"]) else "Company domain"

wb = Workbook()
ws = wb.active
ws.title = "Email list"
ws.append(COLS)
for x in rows:
    ws.append([x.get(c, "") for c in COLS])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F3A5F")
    c.alignment = Alignment(wrap_text=True, vertical="top")
for i, col in enumerate(COLS, 1):
    w = max([len(col)] + [min(len(str(x.get(col, ""))), 55) for x in rows])
    ws.column_dimensions[get_column_letter(i)].width = max(12, min(w + 2, 57))
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions

no = [r for r in d["rows"] if not any(x["Company Name"] == r["Company Name"] for x in rows)]
ws2 = wb.create_sheet("No email found")
cols2 = ["Company Name", "Website", "Country", "Primary Decision Maker", "Primary LinkedIn URL", "LinkedIn Company URL"]
ws2.append(cols2 + ["Note"])
for r in no:
    ws2.append([r.get(c, "") for c in cols2] + ["No email published on site / found via Apify - use contact form or LinkedIn"])
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F3A5F")
for i in range(1, len(cols2) + 2):
    ws2.column_dimensions[get_column_letter(i)].width = 32
wb.save(os.path.join(OUT, "callypso_email_list.xlsx"))
print("emails:", len(rows), "companies:", len({x['Company Name'] for x in rows}), "person:", sum(1 for x in rows if x["Contact Type"].startswith("Decision")),
      "no-email companies:", len(no))
