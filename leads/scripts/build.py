"""Step 7: build final XLSX files + JSON for the web page (deterministic scoring)."""
import json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
import verdicts as V

HERE = os.path.dirname(__file__)
C = os.path.join(HERE, "..", "cache")
OUT = os.path.join(HERE, "..", "..", "reports")
S = json.load(open(os.path.join(C, "sites.json")))
CT = json.load(open(os.path.join(C, "contact.json")))
LI = json.load(open(os.path.join(C, "linkedin_matched.json")))
DM = json.load(open(os.path.join(C, "dm_validated.json")))
DISC = json.load(open(os.path.join(C, "discovered_domains.json")))

CC = {"Germany": "DE", "Netherlands": "NL", "Belgium": "BE", "Austria": "AT", "France": "FR", "United Kingdom": "GB", "Switzerland": "CH"}
COUNTRY_FIX = {"bakkal.eu": ("Estonia", "Tallinn"), "istanbul-market.com": ("France", "Mantes-la-Jolie"),
               "hepsikapinda-shop.com": ("France", "Saint-Dié"), "bellaneva.com": ("Austria", "Wien"),
               "edensmarket.be": ("Belgium", "Brussels"), "thefurniturevilla.co.uk": ("United Kingdom", "London"),
               "mobilyum.co.uk": ("United Kingdom", "London"), "bonvila.com": ("United Kingdom", "London"),
               "aytacfood.co.uk": ("United Kingdom", "London"), "anthap.co.uk": ("United Kingdom", "London")}
EXTRA_CITY = {d: v.get("locations") for d, v in LI.items()}
COLS = ["Company Name", "Website", "Country", "City", "Industry", "Company Size", "E-commerce Confirmed", "E-commerce Platform",
        "Turkish Connection", "Turkish Connection Evidence", "Turkish Evidence URL", "LinkedIn Company URL", "Engage Fit", "Priority",
        "Primary Decision Maker", "Primary Job Title", "Primary LinkedIn URL", "Secondary Decision Maker", "Secondary Job Title",
        "Secondary LinkedIn URL", "Main Source URL", "Confidence Score", "Notes"]


def city_for(d, country):
    if d in COUNTRY_FIX:
        return COUNTRY_FIX[d]
    x = CT.get(d, {})
    if x.get("city") and x.get("city_cc") == CC.get(country):
        return country, x["city"].split()[0].strip(",.")
    for loc in (EXTRA_CITY.get(d) or []):
        if loc.get("country") == CC.get(country) and loc.get("city"):
            return country, loc["city"]
    return country, ""


def size(d):
    r = (LI.get(d) or {}).get("employeeCountRange") or {}
    if r.get("start"):
        return f"{r['start']}-{r['end']} employees (LinkedIn)" if r.get("end") else f"{r['start']}+ employees (LinkedIn)"
    return "Unknown"


def engage_score(d, fit):
    s = S.get(d, {})
    pts = {"H": 6, "M": 3, "L": 0}[fit]
    pts += sum(1 for k in ("newsletter", "account", "wishlist", "discount") if s.get(k))
    pts += 2 if s.get("n_prod_links", 0) >= 30 else 1 if s.get("n_prod_links", 0) >= 10 else 0
    pts += 1 if s.get("platform", "Unknown") != "Unknown" else 0
    pts += 2 if d in DM else 0
    pts += 1 if d in LI else 0
    return pts


def main():
    rows = []
    for d, (name, country, ind, t, ev, fit, note) in V.V.items():
        s = S.get(d, {})
        country, city = city_for(d, country)
        dms = DM.get(d, [])
        li = LI.get(d, {})
        notes = [n for n in [note] if n]
        if li:
            notes.append(f"LinkedIn page matched by {li['match']}")
        if not dms:
            notes.append("No verified decision maker found on LinkedIn")
        plat = s.get("platform", "Unknown").replace("Other: Next.js custom", "Other (custom)").replace("Other: ", "Other – ")
        src = (DISC.get(d, {}).get("urls") or [s.get("final")])[0]
        rows.append({
            "Company Name": name, "Website": s.get("final") or f"https://{d}", "Country": country, "City": city, "Industry": ind,
            "Company Size": size(d), "E-commerce Confirmed": "Yes (cart/checkout + product pages detected)",
            "E-commerce Platform": plat, "Turkish Connection": "Verified (official/reliable source)" if t == "H" else "Strong but limited evidence",
            "Turkish Connection Evidence": ev, "Turkish Evidence URL": s.get("self_turk_url") or s.get("about_url") or s.get("final") or f"https://{d}",
            "LinkedIn Company URL": li.get("linkedinUrl", ""), "Engage Fit": {"H": "HIGH", "M": "MEDIUM", "L": "LOW"}[fit], "Priority": "NORMAL",
            "Primary Decision Maker": dms[0]["name"] if dms else "", "Primary Job Title": dms[0]["title"] if dms else "",
            "Primary LinkedIn URL": dms[0]["url"] if dms else "", "Secondary Decision Maker": dms[1]["name"] if len(dms) > 1 else "",
            "Secondary Job Title": dms[1]["title"] if len(dms) > 1 else "", "Secondary LinkedIn URL": dms[1]["url"] if len(dms) > 1 else "",
            "Main Source URL": src, "Confidence Score": "HIGH" if t == "H" else "MEDIUM", "Notes": "; ".join(notes), "_score": engage_score(d, fit)})
    # Priority HIGH = top third by deterministic Engage score (target was 50 of 150)
    n_high = round(len(rows) / 3)
    for r in sorted(rows, key=lambda r: -r["_score"])[:n_high]:
        r["Priority"] = "HIGH"
    rows.sort(key=lambda r: (r["Priority"] != "HIGH", -r["_score"], r["Country"], r["Company Name"]))

    def sheet(ws, cols, data):
        ws.append(cols)
        for r in data:
            ws.append([r.get(c, "") for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, col in enumerate(cols, 1):
            w = max([len(str(col))] + [min(len(str(r.get(col, ""))), 60) for r in data])
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(w + 2, 62))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    sheet(wb.active, COLS, rows)
    wb.active.title = "Leads"
    wb.save(os.path.join(OUT, "callypso_engage_europe_turkish_ecommerce_leads.xlsx"))

    rev = []
    for d, why in V.REVIEW.items():
        s = S.get(d, {})
        rev.append({"Company / Domain": d, "Website": s.get("final") or f"https://{d}", "Title": s.get("title", ""),
                    "E-commerce detected": "Yes" if s.get("cart") else "Unclear", "Platform": s.get("platform", "Unknown"),
                    "Reason for review": why, "Confidence Score": "LOW", "Main Source URL": (DISC.get(d, {}).get("urls") or [""])[0]})
    big = [{"Company / Domain": d, "Website": (S.get(d, {}).get("final") or f"https://{d}"),
            "Reason for review": "Excluded on request: very large / well-known brand (Turkish connection verified)"} for d in V.EXCLUDED_BIG]
    wb = Workbook()
    rc = ["Company / Domain", "Website", "Title", "E-commerce detected", "Platform", "Reason for review", "Confidence Score", "Main Source URL"]
    sheet(wb.active, rc, rev)
    wb.active.title = "Needs review"
    sheet(wb.create_sheet("Excluded large brands"), ["Company / Domain", "Website", "Reason for review"], big)
    wb.save(os.path.join(OUT, "needs_manual_review.xlsx"))

    stats = {"candidates": len(DISC), "crawled_ok": sum(1 for v in S.values() if v.get("ok")),
             "shoplike": sum(1 for v in S.values() if v.get("ok") and v.get("cart") and (v.get("product_markup") or v.get("n_prod_links", 0) > 3)),
             "verified": len(rows), "high_conf": sum(r["Confidence Score"] == "HIGH" for r in rows),
             "medium_conf": sum(r["Confidence Score"] == "MEDIUM" for r in rows), "high_priority": n_high,
             "with_dm": sum(1 for r in rows if r["Primary Decision Maker"]), "with_li": sum(1 for r in rows if r["LinkedIn Company URL"]),
             "review": len(rev), "excluded_big": len(big)}
    json.dump({"stats": stats, "rows": [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows], "review": rev},
              open(os.path.join(OUT, "leads-data.json"), "w"), ensure_ascii=False)
    print(json.dumps(stats))


if __name__ == "__main__":
    main()
